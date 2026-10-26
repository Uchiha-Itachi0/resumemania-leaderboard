"""Write the ten fictional sample resumes as PDF (and two as DOCX) for demos.

    python scripts/make_sample_resumes.py samples/resumes
"""

from __future__ import annotations

import io
import sys
from html import escape
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from sample_resumes import RESUMES  # noqa: E402

PAGE_MARGIN = 46
CSS = """
* { font-family: sans-serif; }
h1 { font-size: 20px; margin: 0; }
.headline { font-size: 11px; color: #444; margin: 2px 0 0 0; }
.contact { font-size: 9.5px; color: #444; margin: 2px 0 8px 0; }
h2 { font-size: 11px; margin: 10px 0 3px 0; border-bottom: 1px solid #999; text-transform: uppercase; }
p, li { font-size: 9.5px; margin: 1px 0; }
.role { font-weight: bold; font-size: 10px; margin-top: 5px; }
.meta { color: #555; font-size: 9px; }
"""


def _html(r: dict) -> str:
    def section(title: str, body: str) -> str:
        return f"<h2>{escape(title)}</h2>{body}"

    def entries(items: list[tuple]) -> str:
        out = []
        for title, org, dates, bullets in items:
            out.append(f"<p class='role'>{escape(title)} · {escape(org)}</p><p class='meta'>{escape(dates)}</p>")
            out.extend(f"<p>• {escape(b)}</p>" for b in bullets)
        return "".join(out)

    projects = "".join(
        f"<p class='role'>{escape(name)}</p><p class='meta'>{escape(link)}</p>"
        + "".join(f"<p>• {escape(b)}</p>" for b in bullets)
        for name, link, bullets in r["projects"]
    )
    skills = "".join(f"<p><b>{escape(k)}:</b> {escape(v)}</p>" for k, v in r["skills"].items())
    education = "".join(
        f"<p class='role'>{escape(deg)} · {escape(school)}</p><p class='meta'>{escape(dates)}"
        f"{' · ' + escape(extra) if extra else ''}</p>"
        for deg, school, dates, extra in r["education"]
    )
    extras = "".join(f"<p>• {escape(x)}</p>" for x in r["extras"])
    return (
        f"<style>{CSS}</style><h1>{escape(r['name'])}</h1><p class='headline'>{escape(r['headline'])}</p>"
        f"<p class='contact'>{escape(r['contact'])}</p>"
        + section("Summary", f"<p>{escape(r['summary'])}</p>")
        + section("Experience", entries(r["experience"]))
        + section("Projects", projects)
        + section("Skills", skills)
        + section("Education", education)
        + section("Certifications and achievements", extras)
    )


def write_pdf(r: dict, path: Path) -> None:
    import pymupdf

    story = pymupdf.Story(html=_html(r))
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4")
    area = page + (PAGE_MARGIN, PAGE_MARGIN, -PAGE_MARGIN, -PAGE_MARGIN)
    more = 1
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    path.write_bytes(buffer.getvalue())


def write_docx(r: dict, path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(r["name"], level=1)
    d.add_paragraph(r["headline"])
    d.add_paragraph(r["contact"])
    d.add_heading("Summary", level=2)
    d.add_paragraph(r["summary"])
    d.add_heading("Experience", level=2)
    for title, org, dates, bullets in r["experience"]:
        d.add_paragraph(f"{title} · {org} · {dates}").runs[0].bold = True
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Projects", level=2)
    for name, link, bullets in r["projects"]:
        d.add_paragraph(f"{name} · {link}").runs[0].bold = True
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Skills", level=2)
    for k, v in r["skills"].items():
        d.add_paragraph(f"{k}: {v}")
    d.add_heading("Education", level=2)
    for deg, school, dates, extra in r["education"]:
        d.add_paragraph(f"{deg} · {school} · {dates}" + (f" · {extra}" if extra else ""))
    d.add_heading("Certifications and achievements", level=2)
    for x in r["extras"]:
        d.add_paragraph(x, style="List Bullet")
    d.save(path)


def main(out_dir: str) -> None:
    out = Path(out_dir).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    for r in RESUMES:
        if r.get("format") == "docx":
            path = out / f"{r['file']}.docx"
            write_docx(r, path)
        else:
            path = out / f"{r['file']}.pdf"
            write_pdf(r, path)
        print(path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "samples/resumes")
